from pathlib import Path
import struct
from tempfile import TemporaryDirectory
import unittest
import zlib

from docx import Document
from docx.oxml.ns import qn

from md2gdoc.parser import parse_markdown
from md2gdoc.renderer import render_docx


def write_png(path: Path, width: int, height: int) -> Path:
    """Smallest valid PNG of a given size, so tests need no image library."""
    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))

    header = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    rows = b"".join(b"\x00" + b"\xff" * (width * 3) for _ in range(height))
    path.write_bytes(b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) +
                     chunk(b"IDAT", zlib.compress(rows)) + chunk(b"IEND", b""))
    return path


class RendererTests(unittest.TestCase):
    def render(self, markdown):
        directory = TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        output = Path(directory.name) / "document.docx"
        warnings = render_docx(parse_markdown(markdown), output, base_dir=Path(directory.name))
        return Document(output), warnings

    def test_headings_paragraphs_and_unicode(self):
        document, warnings = self.render("# Heading 😀\n\nA paragraph.\n\n## Second")
        self.assertEqual([p.text for p in document.paragraphs], ["Heading 😀", "A paragraph.", "Second"])
        self.assertEqual([p.style.name for p in document.paragraphs], ["Heading 1", "Normal", "Heading 2"])
        self.assertEqual(warnings, ())

    def test_combined_styles_links_and_breaks(self):
        document, _ = self.render("**bold *both*** [**link**](https://example.com/) `code`\nsoft\\\nhard")
        paragraph = document.paragraphs[0]
        self.assertTrue(paragraph.runs[0].bold)
        self.assertTrue(paragraph.runs[1].bold and paragraph.runs[1].italic)
        hyperlink = paragraph.hyperlinks[0]
        self.assertEqual(hyperlink.url, "https://example.com/")
        self.assertTrue(hyperlink.runs[0].bold)
        code = next(run for run in paragraph.runs if run.text == "code")
        self.assertEqual(code.font.name, "Courier New")
        self.assertTrue(code._r.xpath("./w:rPr/w:shd"))
        self.assertIn(" soft\nhard", paragraph.text)

    def test_native_lists_preserve_start_nesting_and_continuations(self):
        document, _ = self.render("3. Three\n4. Four\n   - Nested\n\n     Continuation\n\nBreak\n\n1. Restart")
        paragraphs = document.paragraphs
        self.assertEqual([p.text for p in paragraphs], ["Three", "Four", "Nested", "Continuation", "Break", "Restart"])
        self.assertTrue(paragraphs[0]._p.xpath("./w:pPr/w:numPr"))
        self.assertEqual(paragraphs[2]._p.xpath("./w:pPr/w:numPr/w:ilvl")[0].get(qn("w:val")), "1")
        self.assertFalse(paragraphs[3]._p.xpath("./w:pPr/w:numPr"))
        self.assertNotEqual(paragraphs[0]._p.pPr.numPr.numId.val, paragraphs[5]._p.pPr.numPr.numId.val)
        starts = document.part.numbering_part.element.xpath(".//w:lvl/w:start/@w:val")
        self.assertIn("3", starts)

    def test_code_is_text_with_whitespace_font_and_shading(self):
        code = "\tvalue = 1  \n\n    return value\n"
        document, _ = self.render("```python\n" + code + "```\n")
        paragraph = document.paragraphs[0]
        self.assertEqual(paragraph.text, code)
        self.assertEqual(paragraph.style.name, "Code Block")
        self.assertEqual(paragraph.runs[0].font.name, "Courier New")
        self.assertTrue(paragraph._p.xpath("./w:pPr/w:shd"))
        self.assertFalse(document.inline_shapes)

    def test_tables_keep_alignment_header_and_styles(self):
        document, _ = self.render("| Left | Right |\n| :--- | ---: |\n| **Bold** | [Link](https://example.com/) |")
        table = document.tables[0]
        self.assertEqual(len(table.rows), 2)
        self.assertEqual(len(table.columns), 2)
        self.assertTrue(table.rows[0]._tr.xpath("./w:trPr/w:tblHeader"))
        self.assertTrue(table.cell(1, 0).paragraphs[0].runs[0].bold)
        self.assertEqual(int(table.cell(1, 1).paragraphs[0].alignment), 2)
        self.assertEqual(table.cell(1, 1).paragraphs[0].hyperlinks[0].url, "https://example.com/")

    def test_quote_rule_and_html_are_visible(self):
        document, warnings = self.render("> Quoted\n\n---\n\n<details>Literal HTML</details>\n")
        self.assertEqual(document.paragraphs[0].text, "Quoted")
        self.assertGreater(document.paragraphs[0].paragraph_format.left_indent, 0)
        self.assertTrue(document.paragraphs[1]._p.xpath("./w:pPr/w:pBdr/w:bottom"))
        self.assertIn("<details>", document.paragraphs[2].text)
        self.assertEqual(warnings[0].code, "unsupported_html")

    def test_tall_diagram_fits_the_page_and_warns(self):
        directory = TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        root = Path(directory.name)
        output = root / "document.docx"
        warnings = render_docx(parse_markdown("```mermaid\ngraph TD\n A --> B\n```\n"), output,
                               base_dir=root,
                               mermaid_renderer=lambda source, path: write_png(path, 400, 6000))
        document = Document(output)
        section = document.sections[0]
        shape = document.inline_shapes[0]
        self.assertLessEqual(shape.height, section.page_height - section.top_margin - section.bottom_margin)
        self.assertLessEqual(shape.width, section.page_width - section.left_margin - section.right_margin)
        self.assertEqual(round(shape.width / shape.height, 2), round(400 / 6000, 2))
        self.assertEqual([warning.code for warning in warnings], ["mermaid_scaled_to_page"])

    def test_missing_image_remains_visible_with_warning(self):
        document, warnings = self.render("![A chart](missing.png)")
        self.assertIn("A chart", document.paragraphs[0].text)
        self.assertEqual(warnings[0].code, "image_unavailable")


if __name__ == "__main__":
    unittest.main()

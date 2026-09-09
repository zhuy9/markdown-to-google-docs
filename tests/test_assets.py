from pathlib import Path
import shutil
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import Mock, patch

from docx import Document

from md2gdoc.images import resolve_image
from md2gdoc.parser import parse_markdown
from md2gdoc.renderer import render_docx


FIXTURES = Path(__file__).parent / "fixtures"
PNG = FIXTURES / "assets/example.png"


class AssetTests(unittest.TestCase):
    def setUp(self):
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.directory = Path(temporary.name)

    def test_relative_images_are_embedded_with_alt_links_and_scaled_width(self):
        output = self.directory / "image.docx"
        warnings = render_docx(parse_markdown('[![A chart](assets/example.png "Title")](https://example.com/)'), output, base_dir=FIXTURES)
        document = Document(output)
        self.assertEqual(warnings, ())
        self.assertEqual(len(document.inline_shapes), 1)
        shape = document.inline_shapes[0]
        self.assertEqual(shape._inline.docPr.get("descr"), "A chart")
        self.assertEqual(shape._inline.docPr.get("title"), "Title")
        self.assertTrue(shape._inline.docPr.xpath("./a:hlinkClick"))
        section = document.sections[0]
        self.assertLessEqual(shape.width, section.page_width - section.left_margin - section.right_margin)

    def test_encoded_local_filename(self):
        shutil.copyfile(PNG, self.directory / "a chart.png")
        self.assertEqual(resolve_image("a%20chart.png", self.directory, self.directory), self.directory / "a chart.png")

    def test_remote_images_require_explicit_opt_in(self):
        with patch("md2gdoc.images.urlopen") as download:
            with self.assertRaisesRegex(ValueError, "allow-remote-images"):
                resolve_image("https://example.com/image.png", self.directory, self.directory)
            download.assert_not_called()
            download.return_value.__enter__.return_value.read.return_value = PNG.read_bytes()
            path = resolve_image("https://example.com/image.png", self.directory, self.directory, True)
            self.assertEqual(path.read_bytes(), PNG.read_bytes())

    def test_mermaid_provider_embeds_png_and_retains_source(self):
        output = self.directory / "diagram.docx"
        provider = Mock(side_effect=lambda source, path: Path(shutil.copyfile(PNG, path)))
        warnings = render_docx(parse_markdown("```mermaid\ngraph LR\n A --> B\n```"), output, base_dir=FIXTURES, mermaid_renderer=provider)
        self.assertEqual(warnings, ())
        self.assertEqual(len(Document(output).inline_shapes), 1)
        source = next(output.with_suffix(".assets").glob("*.mmd"))
        self.assertEqual(source.read_text(), "graph LR\n A --> B\n")
        provider.assert_called_once()

    def test_mermaid_failure_retains_selectable_source(self):
        output = self.directory / "failed.docx"
        warnings = render_docx(parse_markdown("```mermaid\ninvalid diagram\n```"), output, base_dir=FIXTURES,
                               mermaid_renderer=Mock(side_effect=RuntimeError("Syntax error")))
        self.assertEqual(warnings[0].code, "mermaid_failed")
        self.assertIn("invalid diagram", Document(output).paragraphs[0].text)
        self.assertTrue(list(output.with_suffix(".assets").glob("*.mmd")))


if __name__ == "__main__":
    unittest.main()

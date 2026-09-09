from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from docx import Document

from md2gdoc.parser import parse_markdown
from md2gdoc.renderer import render_docx
from md2gdoc.validator import validate_docx, validate_google


class ValidationTests(unittest.TestCase):
    def test_google_cannot_hide_missing_paragraph_in_another_word(self):
        remote = {"body": {"content": [{"paragraph": {"elements": [
            {"textRun": {"content": "concatenate\n"}}]}}]}}
        self.assertFalse(validate_google(parse_markdown("cat\n\nconcatenate"), remote)["ok"])

    def test_google_detects_quote_bold_heading_and_code_block_loss(self):
        def remote(text, style=None, paragraph_style=None):
            return {"body": {"content": [{"paragraph": {
                "elements": [{"textRun": {"content": text, "textStyle": style or {}}}],
                "paragraphStyle": paragraph_style or {},
            }}]}}
        self.assertFalse(validate_google(parse_markdown("> **Important**"), remote("Important\n"))["ok"])
        self.assertFalse(validate_google(parse_markdown("## Title"), remote(
            "Title\n", paragraph_style={"namedStyleType": "HEADING_1"}))["ok"])
        mono = {"weightedFontFamily": {"fontFamily": "Courier New"},
                "backgroundColor": {"color": {"rgbColor": {"red": 0.95}}}}
        self.assertFalse(validate_google(parse_markdown("```\na\n```\n\n```\nb\n```"),
                                         remote("a\nb\n", mono))["ok"])

    def test_docx_detects_lost_bold_quote_and_reordered_paragraphs(self):
        source = parse_markdown("> **Important**\n\nFirst\n\nSecond")
        output = self.directory / "styles.docx"
        render_docx(source, output, base_dir=self.directory)
        document = Document(output)
        document.paragraphs[0].runs[0].bold = False
        document.paragraphs[0].paragraph_format.left_indent = 0
        document.paragraphs[1]._p.addprevious(document.paragraphs[2]._p)
        document.save(output)
        report = validate_docx(source, output)
        self.assertFalse(report["ok"])

    def test_google_checks_list_start_nesting_and_table_shape(self):
        source = parse_markdown("3. Item")
        paragraph = {"elements": [{"textRun": {"content": "Item\n"}}],
                     "paragraphStyle": {"indentStart": {"magnitude": 18}},
                     "bullet": {"listId": "list"}}
        level = {"glyphType": "DECIMAL", "startNumber": 3}
        remote = {"body": {"content": [{"paragraph": paragraph}]},
                  "lists": {"list": {"listProperties": {"nestingLevels": [level]}}}}
        self.assertTrue(validate_google(source, remote)["ok"])
        level["startNumber"] = 1
        self.assertFalse(validate_google(source, remote)["ok"])
        level["startNumber"] = 3
        paragraph["bullet"]["nestingLevel"] = 1
        self.assertFalse(validate_google(source, remote)["ok"])
        table_source = parse_markdown("| A | B |\n| --- | --- |")
        cells = [{"content": [{"paragraph": {"elements": [{"textRun": {"content": text + "\n"}}]}}]}
                 for text in ("A", "B")]
        remote = {"body": {"content": [{"table": {"tableRows": [{"tableCells": cells}]}}]}}
        self.assertTrue(validate_google(table_source, remote)["ok"])
        remote["body"]["content"][0]["table"]["tableRows"] = [{"tableCells": [cell]} for cell in cells]
        self.assertFalse(validate_google(table_source, remote)["checks"]["table_shape"])

    def setUp(self):
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.directory = Path(temporary.name)

    def test_valid_document_passes_and_literal_markdown_in_code_is_not_leakage(self):
        source = parse_markdown('# Title\n\n[Link](https://example.com/)\n\n```text\n![literal](image.png)\n```')
        output = self.directory / "valid.docx"
        render_docx(source, output, base_dir=self.directory)
        report = validate_docx(source, output)
        self.assertTrue(report["ok"], report)
        self.assertEqual(report["scope"], "local-docx")
        self.assertTrue(report["checks"]["code_formatting"])

    def test_missing_repeated_paragraph_is_detected(self):
        source = parse_markdown("Same paragraph\n\nSame paragraph")
        output = self.directory / "missing.docx"
        render_docx(source, output, base_dir=self.directory)
        document = Document(output)
        paragraph = document.paragraphs[-1]._p
        paragraph.getparent().remove(paragraph)
        document.save(output)
        self.assertFalse(validate_docx(source, output)["checks"]["text_preserved"])

    def test_lost_code_shading_is_detected(self):
        source = parse_markdown("```\ncode\n```")
        output = self.directory / "code.docx"
        render_docx(source, output, base_dir=self.directory)
        document = Document(output)
        shading = document.paragraphs[0]._p.xpath("./w:pPr/w:shd")[0]
        shading.getparent().remove(shading)
        document.save(output)
        self.assertFalse(validate_docx(source, output)["checks"]["code_formatting"])

    def test_missing_images_and_failed_mermaid_are_not_reported_as_success(self):
        source = parse_markdown("![Missing](missing.png)")
        output = self.directory / "image.docx"
        render_docx(source, output, base_dir=self.directory)
        report = validate_docx(source, output)
        self.assertFalse(report["ok"])
        self.assertEqual(report["counts"]["images"], {"expected": 1, "actual": 0})

    def test_google_readback_checks_actual_heading_and_links(self):
        source = parse_markdown("# Title\n\n[Link](https://example.com/)")
        body = {"content": [
            {"paragraph": {"paragraphStyle": {"namedStyleType": "HEADING_1"}, "elements": [{"textRun": {"content": "Title\n"}}]}},
            {"paragraph": {"elements": [{"textRun": {"content": "Link\n", "textStyle": {"link": {"url": "https://example.com/"}}}}]}},
        ]}
        document = {"tabs": [{"documentTab": {"body": body}}]}
        self.assertTrue(validate_google(source, document)["ok"])
        body["content"][0]["paragraph"]["paragraphStyle"] = {}
        self.assertFalse(validate_google(source, document)["ok"])

    def test_google_connector_flattened_tabs_and_native_child_tabs(self):
        source = parse_markdown("# First\n\n# Second")
        bodies = [{"content": [{"paragraph": {
            "paragraphStyle": {"namedStyleType": "HEADING_1"},
            "elements": [{"textRun": {"content": text + "\n"}}],
        }}]} for text in ("First", "Second")]
        for tabs in (
            [{"tabId": "first", "body": bodies[0]}, {"tabId": "second", "body": bodies[1]}],
            [{"documentTab": {"body": bodies[0]}, "childTabs": [{"documentTab": {"body": bodies[1]}}]}],
        ):
            with self.subTest(tabs=tabs):
                self.assertTrue(validate_google(source, {"tabs": tabs})["ok"])

    def test_google_imported_code_keeps_soft_breaks_indentation_and_paragraph_shading(self):
        source = parse_markdown("```python\ndef identity(value):\n    return value\n```")
        background = {"color": {"rgbColor": {"red": 0.95}}}
        run = {"content": "def identity(value):\u000b    return value\u000b\n", "textStyle": {
            "weightedFontFamily": {"fontFamily": "Courier New"}, "backgroundColor": {},
        }}
        paragraph = {"paragraphStyle": {"shading": {"backgroundColor": background}},
                     "elements": [{"textRun": run}]}
        document = {"body": {"content": [{"paragraph": paragraph}]}}
        self.assertTrue(validate_google(source, document)["ok"])
        paragraph["paragraphStyle"]["shading"]["backgroundColor"] = {}
        self.assertFalse(validate_google(source, document)["checks"]["code_formatting"])
        run["textStyle"]["backgroundColor"] = background
        self.assertTrue(validate_google(source, document)["ok"])
        run["content"] = run["content"].replace("    return", "return")
        self.assertFalse(validate_google(source, document)["ok"])

    def test_google_zero_width_border_is_not_a_horizontal_rule(self):
        source = parse_markdown("Paragraph\n\n---")
        document = {"body": {"content": [
            {"paragraph": {"paragraphStyle": {"borderBottom": {"width": {"unit": "PT"}}},
                           "elements": [{"textRun": {"content": "Paragraph\n"}}]}},
            {"paragraph": {"paragraphStyle": {"borderBottom": {"width": {"magnitude": 0.75, "unit": "PT"}}}}},
        ]}}
        self.assertTrue(validate_google(source, document)["ok"])
        document["body"]["content"][1]["paragraph"]["paragraphStyle"]["borderBottom"]["width"]["magnitude"] = 0
        self.assertFalse(validate_google(source, document)["ok"])


if __name__ == "__main__":
    unittest.main()

"""Render the shared IR as editable DOCX, including embedded diagram images."""

from collections.abc import Callable
from hashlib import sha256
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.exceptions import InvalidImageStreamError, UnexpectedEndOfFileError, UnrecognizedImageError
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor

from . import models as ir
from .images import resolve_image
from .mermaid import render_mermaid


def _xml(tag: str, **attributes):
    element = OxmlElement(tag)
    for key, value in attributes.items():
        element.set(qn("w:" + key), str(value))
    return element


def render_docx(
    document: ir.Document,
    output: Path,
    *,
    base_dir: Path,
    mermaid_renderer: Callable[[str, Path], Path] = render_mermaid,
    allow_remote_images: bool = False,
) -> tuple[ir.Warning, ...]:
    writer = _Writer(document, output, base_dir, mermaid_renderer, allow_remote_images)
    for block in document.blocks:
        writer.block(block)
    output.parent.mkdir(parents=True, exist_ok=True)
    writer.doc.save(output)
    return tuple(writer.warnings)


class _Writer:
    def __init__(self, document, output, base_dir, mermaid_renderer, allow_remote_images):
        self.doc = Document()
        self.doc.core_properties.author = ""
        self.doc.core_properties.last_modified_by = ""
        self.doc.core_properties.comments = ""
        self.doc.core_properties.title = output.stem
        self.doc.styles["Normal"].font.name = "Calibri"
        self.doc.styles["Normal"].font.size = Pt(11)
        self.doc.styles["Normal"].paragraph_format.space_after = Pt(6)
        code = self.doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code.font.name = "Courier New"
        code.font.size = Pt(10)
        code.paragraph_format.space_after = Pt(6)
        self.warnings = list(document.warnings)
        self.base_dir = base_dir
        self.assets = output.with_suffix(".assets")
        self.mermaid_renderer = mermaid_renderer
        self.allow_remote_images = allow_remote_images
        section = self.doc.sections[0]
        self.width = section.page_width - section.left_margin - section.right_margin

    def paragraph(self, block, depth, existing=None):
        paragraph = existing if existing is not None else self.doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(depth / 4)
        if isinstance(block, ir.Heading):
            paragraph.style = f"Heading {block.level}"
        self.inlines(paragraph, block.inlines, block.source, self.width - Inches(depth / 4))

    def block(self, block: ir.Block, depth: int = 0):
        if isinstance(block, (ir.Paragraph, ir.Heading)):
            self.paragraph(block, depth)
        elif isinstance(block, ir.ListBlock):
            self.list_block(block, depth)
        elif isinstance(block, ir.Blockquote):
            for child in block.blocks:
                self.block(child, depth + 1)
        elif isinstance(block, ir.CodeBlock):
            paragraph = self.doc.add_paragraph(style="Code Block")
            paragraph.paragraph_format.left_indent = Inches(depth / 4)
            paragraph._p.get_or_add_pPr().append(_xml("w:shd", fill="F3F4F6", val="clear"))
            run = paragraph.add_run(block.text)
            run.font.name = "Courier New"
            run.font.size = Pt(10)
        elif isinstance(block, ir.Table):
            self.table(block)
        elif isinstance(block, ir.MermaidDiagram):
            self.diagram(block, depth)
        elif isinstance(block, ir.HorizontalRule):
            borders = _xml("w:pBdr")
            borders.append(_xml("w:bottom", val="single", sz=6, color="B8BDC5", space=1))
            self.doc.add_paragraph()._p.get_or_add_pPr().append(borders)

    def inlines(self, paragraph, inlines, source, width, header=False):
        for inline in inlines:
            if isinstance(inline, ir.LineBreak):
                paragraph.add_run("\n" if inline.hard else " ")
            elif isinstance(inline, ir.Image):
                self.image(paragraph, inline, source, width)
            else:
                run = paragraph.add_run(inline.text)
                run.bold = inline.bold or header or None
                run.italic = inline.italic or None
                if inline.code:
                    run.font.name = "Courier New"
                    run.font.size = Pt(10)
                    run._r.get_or_add_rPr().append(_xml("w:shd", fill="F3F4F6", val="clear"))
                if inline.link:
                    run.font.color.rgb = RGBColor.from_string("0563C1")
                    run.underline = True
                    hyperlink = OxmlElement("w:hyperlink")
                    hyperlink.set(qn("r:id"), paragraph.part.relate_to(inline.link, RT.HYPERLINK, is_external=True))
                    hyperlink.append(run._r)
                    paragraph._p.append(hyperlink)

    def numbering(self, block, depth):
        root = self.doc.part.numbering_part.element
        abstracts = root.findall(qn("w:abstractNum"))
        identifier = max((int(item.get(qn("w:abstractNumId"))) for item in abstracts), default=-1) + 1
        abstract = _xml("w:abstractNum", abstractNumId=identifier)
        abstract.append(_xml("w:multiLevelType", val="multilevel"))
        for level in range(9):
            entry = _xml("w:lvl", ilvl=level)
            entry.append(_xml("w:start", val=block.start if level == min(depth, 8) else 1))
            entry.append(_xml("w:numFmt", val="decimal" if block.ordered else "bullet"))
            entry.append(_xml("w:lvlText", val=f"%{level + 1}." if block.ordered else "•"))
            properties = _xml("w:pPr")
            properties.append(_xml("w:ind", left=360 * (level + 1), hanging=360))
            entry.append(properties)
            abstract.append(entry)
        root.insert(len(abstracts), abstract)
        return root.add_num(identifier).numId

    def list_block(self, block, depth):
        number = self.numbering(block, depth)
        if depth > 8:
            self.warnings.append(ir.Warning("list_depth", "DOCX supports nine numbering levels; deeper items retain indentation.", block.source))
        for item in block.items:
            marker = self.doc.add_paragraph()
            marker.paragraph_format.left_indent = Inches((depth + 1) / 4)
            properties = marker._p.get_or_add_pPr().get_or_add_numPr()
            properties.get_or_add_ilvl().val = min(depth, 8)
            properties.get_or_add_numId().val = number
            for index, child in enumerate(item.blocks):
                if index == 0 and isinstance(child, (ir.Paragraph, ir.Heading)):
                    self.paragraph(child, depth + 1, marker)
                else:
                    self.block(child, depth + 1)

    def table(self, block):
        table = self.doc.add_table(rows=1 + len(block.rows), cols=len(block.header), style="Table Grid")
        table.autofit = False
        table.rows[0]._tr.get_or_add_trPr().append(_xml("w:tblHeader"))
        alignments = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}
        for row_index, cells in enumerate((block.header, *block.rows)):
            for column, cell in enumerate(cells):
                target = table.cell(row_index, column)
                paragraph = target.paragraphs[0]
                paragraph.alignment = alignments.get(block.alignments[column])
                self.inlines(paragraph, cell.inlines, block.source, target.width, header=row_index == 0)
                if row_index == 0:
                    target._tc.get_or_add_tcPr().append(_xml("w:shd", fill="EEF2F6", val="clear"))

    def picture(self, paragraph, path, alt, title, width):
        shape = paragraph.add_run().add_picture(str(path))
        if shape.width > width:
            shape.height = round(shape.height * width / shape.width)
            shape.width = width
        shape._inline.docPr.set("descr", alt)
        if title:
            shape._inline.docPr.set("title", title)
        return shape

    def image(self, paragraph, image, source, width):
        try:
            path = resolve_image(image.src, self.base_dir, self.assets, self.allow_remote_images)
            shape = self.picture(paragraph, path, image.alt, image.title, width)
            if image.link:
                link = OxmlElement("a:hlinkClick")
                link.set(qn("r:id"), paragraph.part.relate_to(image.link, RT.HYPERLINK, is_external=True))
                shape._inline.docPr.append(link)
        except (OSError, ValueError, InvalidImageStreamError, UnexpectedEndOfFileError, UnrecognizedImageError) as error:
            paragraph.add_run(f"[Image unavailable: {image.alt or image.src}]")
            self.warnings.append(ir.Warning("image_unavailable", str(error), source))

    def diagram(self, block, depth):
        name = "mermaid-" + sha256(block.text.encode()).hexdigest()[:16]
        path = self.assets / (name + ".png")
        self.assets.mkdir(parents=True, exist_ok=True)
        path.with_suffix(".mmd").write_text(block.text, encoding="utf-8")
        try:
            rendered = self.mermaid_renderer(block.text, path)
            paragraph = self.doc.add_paragraph()
            shape = self.picture(paragraph, rendered, "Mermaid diagram", None, self.width - Inches(depth / 4))
            shape._inline.docPr.set("name", name)
        except (OSError, RuntimeError, ValueError, InvalidImageStreamError, UnexpectedEndOfFileError, UnrecognizedImageError) as error:
            self.warnings.append(ir.Warning("mermaid_failed", str(error), block.source))
            self.block(ir.CodeBlock(block.text, "mermaid", block.source), depth)
